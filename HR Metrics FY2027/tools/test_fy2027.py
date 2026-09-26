"""End-to-end test: fill random inputs, recalc with LibreOffice, verify against Python-computed expectations,
then flatten the external links of 0 GCG and GEL with the recalculated source values and verify the roll-ups."""
import os, sys, json, random, shutil, subprocess, datetime, re, glob
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_fy2027 as B

RECALC = '/root/.claude/skills/synced/97121f19-48cb-420b-b2d5-a5562825e984_42d4a032-2d33-4cbb-a129-fd74d6e8c0f3/xlsx/scripts/recalc.py'
OUT = os.path.join(HERE, 'out')
T = os.path.join(HERE, 'test')
random.seed(7)
fails = []

def check(cond, msg):
    if not cond:
        fails.append(msg)

def close(a, b, tol=1e-6):
    if a in (None, '') and b in (None, ''):
        return True
    try:
        return abs(float(a) - float(b)) <= tol * max(1, abs(float(b)))
    except Exception:
        return False

def recalc(path):
    r = subprocess.run([sys.executable, RECALC, path, '600'] + (['--force'] if 'flat' in path else []),
                       capture_output=True, text=True)
    res = json.loads(r.stdout)
    check(res.get('status') == 'success', f'recalc {os.path.basename(path)}: {json.dumps(res)[:600]}')
    return res

INPUTS = [r for r in B.INPUT_ROWS if r not in B.TAB_ROWS]

def fill_entity(path):
    wb = openpyxl.load_workbook(path)
    comps = [(c, wb['OCT'].cell(2, c).value) for c in range(3, wb['OCT'].max_column + 1)
             if wb['OCT'].cell(2, c).value not in (None, 'TOTAL', 'Total by Division') and
             not str(wb['OCT'].cell(2, c).value).startswith('=')]
    comps = [(L(c), n) for c, n in comps]
    tabs = wb.sheetnames[13:]
    for mi, m in enumerate(B.MONTHS):
        ws = wb[m]
        for k, (col, _) in enumerate(comps):
            if mi >= 10 and k == 0:          # first company has not reported AUG/SEP yet
                continue
            hc_perm, hc_temp = random.randint(20, 300), random.randint(0, 40)
            vals = {3: hc_perm, 4: hc_temp, 6: hc_perm + hc_temp - 7, 7: 0 if k % 2 else None, 8: 4, 9: 3,
                    10: 173.33, 12: random.randint(0, 30), 13: random.randint(0, 3000), 18: round(random.uniform(1, 5), 2)}
            for r in INPUTS:
                if r in vals:
                    v = vals[r]
                elif r == 41:
                    v = round(random.uniform(0.3, 1), 3)
                elif r in (58, 59):
                    v = round(random.uniform(0, 5000), 2)
                else:
                    v = random.randint(0, 25)
                ws[f'{col}{r}'].value = v
            if mi == 0:
                ws[f'{col}31'].value = hc_perm + hc_temp + random.randint(-5, 5)
    for t in tabs:
        ws = wb[t]
        for r in range(4, 16):
            ws[f'D{r}'].value = random.choice([20, 21, 22, 23])
        # permanent + temporary hires: some before FY2027 (excluded), some in FY2027
        for i in range(6):
            r = 21 + i
            od = datetime.datetime(2026, 8, 1) + datetime.timedelta(days=random.randint(0, 200))
            sd = od + datetime.timedelta(days=random.randint(5, 90))
            ws[f'B{r}'].value = f'Job {i}'; ws[f'C{r}'].value = random.choice([1, 2, None])
            ws[f'D{r}'].value = od; ws[f'E{r}'].value = sd
            ws[f'I{r}'].value = f'Temp {i}'; ws[f'J{r}'].value = random.choice([1, 3])
            ws[f'K{r}'].value = od; ws[f'L{r}'].value = sd + datetime.timedelta(days=3)
        ws['B27'].value = 'Open position'; ws['C27'].value = 1; ws['D27'].value = datetime.datetime(2026, 11, 3)
    wb.save(path)
    return comps, tabs

def expected_entity(path, comps, tabs):
    """Independent Python computation for a handful of key metrics."""
    wb = openpyxl.load_workbook(path, data_only=True)   # recalculated values
    wi = wb  # inputs are plain values in the same file
    tot = L(CI(comps[-1][0]) + 1)
    for mi, m in enumerate(B.MONTHS):
        ws = wb[m]
        hc_sum, rate_num, rate_den = 0, 0, 0
        any_hc = False
        for col, name in comps:
            g = lambda r: ws[f'{col}{r}'].value
            if g(3) is None:
                check(g(5) in (None, ''), f'{path}:{m}!{col}5 should be blank')
                check(g(70) == 35, f'{path}:{m}!{col}70 input-not-submitted = {g(70)} (expected 35)')
                continue
            any_hc = True
            hc = g(3) + g(4); hc_sum += hc
            check(close(g(5), hc), f'{m}!{col}5')
            fte = g(6) + (g(7) or 0) * .75 + g(8) * .5 + g(9) * .25
            check(close(g(11), fte), f'{m}!{col}11 {g(11)} vs {fte}')
            last = g(31)
            check(close(g(33), g(24) / ((hc + last) / 2)), f'{m}!{col}33 turnover')
            check(close(g(37), hc + g(36)), f'{m}!{col}37 authorized')
            # company tab absence
            tab = [t for t in tabs if wb[t]['A1'].value == name][0]
            tw = wb[tab]
            wd = tw[f'D{4 + mi}'].value
            exp46 = g(43) / (hc * wd)
            check(close(g(46), exp46), f'{m}!{col}46 absence {g(46)} vs {exp46}')
            check(close(g(48), exp46 + g(44) / (hc * wd)), f'{m}!{col}48 absence total')
            rate_num += g(46) * hc; rate_den += hc
            # time to hire YTD up to month end
            me = B.MONTH_DATES[mi]
            nxt = datetime.datetime(me.year + (me.month == 12), me.month % 12 + 1, 1)
            num = den = 0
            for r in range(21, 27):
                od, sd, p = tw[f'D{r}'].value, tw[f'E{r}'].value, tw[f'C{r}'].value or 1
                if datetime.datetime(2026, 10, 1) <= sd < nxt:
                    num += (sd - od).days * p; den += p
            exp39 = num / den if den else None
            check(close(g(39), exp39), f'{path}:{m}!{col}39 TTH {g(39)} vs {exp39}')
            exp70 = (g(7) is None) + (g(39) in (None, '')) + (g(40) in (None, ''))
            check(g(70) == exp70, f'{path}:{m}!{col}70 = {g(70)} vs {exp70}')
        if any_hc:
            check(close(ws[f'{tot}5'].value, hc_sum), f'{m}!{tot}5 total HC')
            check(close(ws[f'{tot}46'].value, rate_num / rate_den), f'{m}!{tot}46 weighted absence')
    # Quarter
    q = wb['Quarter']
    for m, qc in B.Q_MONTH_COL.items():
        for r in (3, 5, 11, 22, 46, 70):
            check(close(q[f'{qc}{r}'].value, wb[m][f'{tot}{r}'].value), f'Quarter!{qc}{r} vs {m}!{tot}{r}')
    check(close(q['F19'].value, sum(wb[m][f'{tot}19'].value for m in B.MONTHS[:3])), 'Quarter F19 sum')
    check(close(q['F3'].value, sum(wb[m][f'{tot}3'].value for m in B.MONTHS[:3]) / 3), 'Quarter F3 avg')
    latest = [wb[m][f'{tot}39'].value for m in B.MONTHS if wb[m][f'{tot}39'].value not in (None, '')][-1]
    check(close(q['S39'].value, latest), 'Quarter S39 latest YTD')
    q4 = [wb[m][f'{tot}39'].value for m in B.MONTHS[9:] if wb[m][f'{tot}39'].value not in (None, '')][-1]
    check(close(q['R39'].value, q4), 'Quarter R39 latest YTD in Q4')
    return wb

def flatten(src_path, dst_path, sources):
    """Replace [n]Sheet!Cell link formulas by the recalculated value of that cell in sources[n-1]."""
    wb = openpyxl.load_workbook(src_path)
    vals = [openpyxl.load_workbook(p, data_only=True) for p in sources]
    pat = re.compile(r"^=IF\(\[(\d+)\](\w+)!([A-Z]+\d+)=\"\",\"\",\[\1\]\2!\3\)$")
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and '[' in c.value:
                    mt = pat.match(c.value)
                    assert mt, c.value
                    v = vals[int(mt.group(1)) - 1][mt.group(2)][mt.group(3)].value
                    c.value = '' if v is None else v
                    if c.value == '':
                        c.value = '=""'
                    n += 1
    wb._external_links = []
    wb.save(dst_path)
    return n

def main():
    shutil.rmtree(T, ignore_errors=True)
    shutil.copytree(OUT, T)
    entities = {}
    for p in sorted(glob.glob(os.path.join(T, '**', '*.xlsx'), recursive=True)):
        if '0 GCG' in p or 'GEL HR' in p:
            continue
        comps, tabs = fill_entity(p)
        recalc(p)
        expected_entity(p, comps, tabs)
        entities[os.path.basename(p)] = p
        print('verified', os.path.basename(p), len(fails), 'fails so far')

    # 0 GCG: flatten against recalculated clusters
    g_src = os.path.join(T, B.GCG_DIR, B.GCG_FILE)
    order = [os.path.join(T, B.GCG_DIR, f) for _, _, f in B.CLUSTERS]
    g_flat = g_src.replace('.xlsx', ' flat.xlsx')
    print('flattened links:', flatten(g_src, g_flat, order))
    recalc(g_flat)
    gv = openpyxl.load_workbook(g_flat, data_only=True)
    cl = {k: openpyxl.load_workbook(p, data_only=True) for (k, _, _), p in zip(B.CLUSTERS, order)}
    # every company column (all 68 rows, incl. derived rows recomputed locally) must equal the cluster
    comp_pos = {}
    for k, *_ in B.CLUSTERS:
        cols = [c for c, v in B.GCG_MAP.items() if v[0] == k]
        for i, mc in enumerate(cols):
            comp_pos[mc] = (k, L(3 + i))
    for m in B.MONTHS:
        for mc, (k, cc) in comp_pos.items():
            for r in range(3, 71):
                if r == 49:
                    continue   # weight factor is relative to a different total
                a, b = gv[m][f'{mc}{r}'].value, cl[k][m][f'{cc}{r}'].value
                check(close(a, b), f'0 GCG {m}!{mc}{r}={a} vs cluster {k} {cc}{r}={b}')
        for r in (3, 5, 19, 22, 43, 70):
            tot = sum(float(cl[k][m][f'{L(CI(max([v[1] for v in comp_pos.values() if v[0]==k], key=CI))+1)}{r}'].value or 0)
                      for k, *_ in B.CLUSTERS)
            check(close(gv[m][f'AJ{r}'].value, tot), f'0 GCG {m}!AJ{r} {gv[m][f"AJ{r}"].value} vs sum clusters {tot}')
    print('0 GCG verified', len(fails), 'fails so far')

    # GEL: flatten against recalculated divisions + flat 0 GCG
    gel_src = os.path.join(T, B.GEL_DIR, B.GEL_FILE)
    srcs = [os.path.join(T, folder, f) for _, folder, f, _, _ in B.DIVISIONS] + [g_flat]
    gel_flat = gel_src.replace('.xlsx', ' flat.xlsx')
    print('flattened links:', flatten(gel_src, gel_flat, srcs))
    recalc(gel_flat)
    gl = openpyxl.load_workbook(gel_flat, data_only=True)['GEL Divisions']
    qs = [openpyxl.load_workbook(p, data_only=True)['Quarter'] for p in srcs]
    ncol = 5
    for d, qd in enumerate(qs):
        for i, qc in enumerate(B.Q_COLS):
            c = L(3 + d * ncol + i)
            for r in range(3, 71):
                if r != 49:
                    check(close(gl[f'{c}{r}'].value, qd[f'{qc}{r}'].value), f'GEL {c}{r} vs div{d} Quarter!{qc}{r}')
    for i, qc in enumerate(B.Q_COLS):
        g = L(3 + len(qs) * ncol + i)
        for r in (3, 4, 19, 22, 43, 55, 58, 70):
            s = sum(float(qd[f'{qc}{r}'].value or 0) for qd in qs)
            check(close(gl[f'{g}{r}'].value, s), f'GEL {g}{r}={gl[f"{g}{r}"].value} vs {s}')
        check(close(gl[f'{g}5'].value, gl[f'{g}3'].value + gl[f'{g}4'].value), f'GEL {g}5')
        hc = [float(qd[f'{qc}5'].value) for qd in qs]
        ab = [float(qd[f'{qc}46'].value) for qd in qs]
        check(close(gl[f'{g}46'].value, sum(a * h for a, h in zip(ab, hc)) / sum(hc)), f'GEL {g}46 weighted')
        check(close(gl[f'{g}49'].value, 1), f'GEL {g}49 weights sum to 1')
        ex, l_, t_ = (float(gl[f'{g}24'].value), float(gl[f'{g}31'].value), float(gl[f'{g}32'].value))
        check(close(gl[f'{g}33'].value, ex / ((l_ + t_) / 2)), f'GEL {g}33 turnover')
    print('GEL verified')
    print('FAILS:', len(fails))
    for f in fails[:60]:
        print('  ', f)

if __name__ == '__main__':
    main()
